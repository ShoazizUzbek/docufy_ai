from io import BytesIO
from unittest.mock import patch

import pymupdf
from django.test import SimpleTestCase
from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK

from documents.processing.extractors import extract_docx, extract_pdf, extract_txt


class ExtractPdfTests(SimpleTestCase):
    def _build_pdf_bytes(self, build_fn) -> bytes:
        doc = pymupdf.open()
        build_fn(doc)
        data = doc.tobytes()
        doc.close()
        return data

    def test_extracts_text_layer_with_page_numbers_and_headings(self):
        def build(doc):
            page1 = doc.new_page()
            page1.insert_text((72, 72), 'Chapter 2', fontsize=20)
            page1.insert_text((72, 110), 'The tenant shall pay rent on the first of each month.', fontsize=11)
            page2 = doc.new_page()
            page2.insert_text((72, 72), 'This paragraph lives on the second page of the document.', fontsize=11)

        file_bytes = self._build_pdf_bytes(build)
        blocks, page_count, used_ocr = extract_pdf(file_bytes)

        self.assertFalse(used_ocr)
        self.assertEqual(page_count, 2)

        heading_blocks = [b for b in blocks if b.is_heading]
        self.assertTrue(any(b.text.strip() == 'Chapter 2' for b in heading_blocks))

        page_numbers = {b.page_number for b in blocks}
        self.assertEqual(page_numbers, {1, 2})

        second_page_block = next(b for b in blocks if 'second page' in b.text)
        self.assertEqual(second_page_block.page_number, 2)

    @patch('documents.processing.extractors.run_ocr')
    def test_pdf_without_text_layer_falls_back_to_ocr(self, mock_run_ocr):
        mock_run_ocr.return_value = ['CHAPTER 1', 'Recognized OCR text line.']

        def build(doc):
            page = doc.new_page()
            # No inserted text -> no text layer, image-only page.
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100))
            pix.set_rect(pix.irect, (255, 255, 255))
            page.insert_image(page.rect, pixmap=pix)

        file_bytes = self._build_pdf_bytes(build)
        blocks, page_count, used_ocr = extract_pdf(file_bytes)

        self.assertTrue(used_ocr)
        self.assertEqual(page_count, 1)
        self.assertTrue(mock_run_ocr.called)
        self.assertTrue(any('Recognized OCR text line' in b.text for b in blocks))
        self.assertTrue(any(b.is_heading for b in blocks))  # "CHAPTER 1" via keyword heuristic


class ExtractDocxTests(SimpleTestCase):
    def test_extracts_headings_and_tracks_manual_page_breaks(self):
        doc = DocxDocument()
        doc.add_paragraph('Chapter 2', style='Heading 1')
        doc.add_paragraph('Article 14', style='Heading 2')
        doc.add_paragraph('The lessee agrees to maintain the property in good condition.')
        page_break_paragraph = doc.add_paragraph()
        page_break_paragraph.add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph('This paragraph is on the second page.')

        buf = BytesIO()
        doc.save(buf)

        blocks, page_count = extract_docx(buf.getvalue())

        self.assertEqual(page_count, 2)

        chapter_block = next(b for b in blocks if b.text == 'Chapter 2')
        self.assertTrue(chapter_block.is_heading)
        self.assertEqual(chapter_block.heading_level, 1)

        article_block = next(b for b in blocks if b.text == 'Article 14')
        self.assertTrue(article_block.is_heading)
        self.assertEqual(article_block.heading_level, 2)

        second_page_block = next(b for b in blocks if 'second page' in b.text)
        self.assertEqual(second_page_block.page_number, 2)


class ExtractTxtTests(SimpleTestCase):
    def test_splits_paragraphs_and_detects_heading_keywords(self):
        content = (
            'Chapter 1: Introduction\n\n'
            'This agreement is entered into by both parties as of the effective date.\n\n'
            'Article 1\n\n'
            'The parties shall act in good faith at all times.'
        )
        blocks, page_count = extract_txt(content.encode('utf-8'))

        self.assertEqual(page_count, 1)
        self.assertTrue(all(b.page_number == 1 for b in blocks))

        heading_texts = [b.text for b in blocks if b.is_heading]
        self.assertIn('Chapter 1: Introduction', heading_texts)
        self.assertIn('Article 1', heading_texts)

    def test_form_feed_splits_pages(self):
        content = 'First page text.\x0cSecond page text.'
        blocks, page_count = extract_txt(content.encode('utf-8'))

        self.assertEqual(page_count, 2)
        self.assertEqual(blocks[0].page_number, 1)
        self.assertEqual(blocks[1].page_number, 2)
